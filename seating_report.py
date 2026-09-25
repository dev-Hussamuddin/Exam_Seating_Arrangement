"""Excel presentation of committed allocation summaries; no database writes."""

from collections import defaultdict
from datetime import date, datetime, time, timedelta
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import OUTPUT_DIR


def _date(value):
    if isinstance(value, datetime):
        return value.date()
    return date.fromisoformat(value) if isinstance(value, str) else value


def _time(value):
    if isinstance(value, timedelta):
        return (datetime.min + value).time()
    return time.fromisoformat(value) if isinstance(value, str) else value


def _staff_style(sheet, header_row, widths):
    """Shared print styling for the block overview and numbered bench maps."""
    border = Border(*( [Side(style="thin", color="D5DCE4")] * 4))
    for row in sheet:
        for cell in row:
            cell.font = Font(name="Calibri", size=11, color="203040")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, str):
                cell.data_type = "s"
            if cell.row >= header_row:
                cell.border = border
            if cell.row in (1, header_row):
                cell.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="203864")
            elif cell.column == 1 and cell.row < header_row:
                cell.font = Font(name="Calibri", size=11, bold=True)
    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[header_row].height = 34
    for row in range(header_row + 1, sheet.max_row + 1):
        lines = max((len(str(sheet.cell(row, col).value or "")) // max(1, width - 3) + 1)
                    for col, width in enumerate(widths, 1))
        sheet.row_dimensions[row].height = max(30, 15 * lines)
    sheet.freeze_panes = f"D{header_row + 1}"
    last = get_column_letter(len(widths))
    sheet.auto_filter.ref = f"A{header_row}:{last}{sheet.max_row}"
    sheet.print_title_rows = f"1:{header_row}"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if len(widths) > 6 else sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_area = f"A1:{last}{sheet.max_row}"
    sheet.print_options.horizontalCentered = True


def _exam_heading(sheet, slots):
    days = sorted({_date(slot[0]) for slot in slots})
    sheet.append(["Exam Date", days[0] if len(days) == 1 else ", ".join(d.isoformat() for d in days)])
    sheet["B2"].number_format = "dd-mmm-yyyy"
    intervals = {(_time(slot[1]), _time(slot[2])) for slot in slots}
    if len(intervals) == 1:
        start, end = next(iter(intervals))
        sheet.append(["Exam Time", start if start is not None else "Not configured", end])
        sheet["B3"].number_format = sheet["C3"].number_format = "hh:mm"
    else:
        sheet.append(["Exam Time", "See slot headings / room seat maps" if slots else "Not configured"])


def add_block_and_seat_sheets(workbook, summaries):
    from seating_allocator import compact_ranges, display_time, validate_allocation
    from seating_blocks import make_blocks

    overview = workbook.create_sheet("Room Blocks", 0)
    overview.append(["ROOM-WISE EXAM SEATING — BLOCKS"])
    overview.merge_cells("A1:I1")
    _exam_heading(overview, [slot for slot, _, _, _ in summaries])
    overview.append([])
    all_blocks = []
    for slot, exams, rooms, allocations in summaries:
        validate_allocation(exams, rooms, allocations)
        blocks = make_blocks(allocations, len(all_blocks) + 1)
        all_blocks.extend(blocks)
    eligible = sum(len(e["rolls"]) for _, exams, _, _ in summaries for e in exams)
    allocated = sum(len(b["students"]) for b in all_blocks)
    capacity = sum(r["capacity"] for _, _, rooms, _ in summaries for r in rooms)
    for label, value in (("Total Eligible Students", eligible), ("Total Allocated Students", allocated),
                         ("Total Capacity Used", allocated), ("Total Unused Seats", capacity - allocated),
                         ("Total Blocks", len(all_blocks))):
        overview.append([label, value])
    overview.append(["Blocks target 30 students; remainders are absorbed. Mixed subjects share a block number; rows show each subject's allocation."])
    overview.merge_cells("A10:I10")
    overview.row_dimensions[10].height = 30
    overview.append([])
    overview.append(["Room Number", "Room Capacity", "Block Number", "Class", "Department", "Subject",
                     "Allocated Roll Numbers", "Number of Students Allocated", "Unused Seats"])
    next_block, map_number = 1, 0
    for slot, exams, rooms, allocations in summaries:
        blocks = make_blocks(allocations, next_block)
        next_block += len(blocks)
        by_exam = {e["timetable_id"]: e for e in exams}
        if len(summaries) > 1:
            overview.append([f"Slot: {_date(slot[0]).isoformat()} {display_time(slot[1])}–{display_time(slot[2])}"])
            overview.merge_cells(start_row=overview.max_row, start_column=1, end_row=overview.max_row, end_column=9)
        for room in rooms:
            room_blocks = [b for b in blocks if b["classroom_id"] == room["classroom_id"]]
            used = sum(len(b["students"]) for b in room_blocks)
            detail_rows = [(block, group) for block in room_blocks for group in block["groups"]]
            for index, (block, group) in enumerate(detail_rows or [(None, None)]):
                exam = by_exam[group["timetable_id"]] if group else {}
                overview.append([str(room["classroom_no"]), room["capacity"] if index == 0 else None,
                                 block["block_number"] if block else None, exam.get("class_name", ""),
                                 (exam.get("department") or "Not specified") if block else "", exam.get("subject", ""),
                                 compact_ranges(group["rolls"]) if group else "",
                                 len(group["rolls"]) if group else 0, room["capacity"] - used if index == 0 else None])
                for cell in overview[overview.max_row]:
                    cell.fill = PatternFill("solid", fgColor="E8EFF7" if map_number % 2 == 0 else "FFFFFF")
            map_number += 1
            sheet = workbook.create_sheet(f"Seats {map_number:03d}")
            sheet.append([f"ROOM {room['classroom_no']} — NUMBERED BENCH / SEAT MAP"])
            sheet.merge_cells("A1:F1")
            _exam_heading(sheet, [slot])
            sheet.append(["Room Capacity", room["capacity"], "Allocated", used, "Unused Seats", room["capacity"] - used])
            sheet.append(["One student per seat. Seat numbers follow allocation order; physical room rows/columns are not specified."])
            sheet.merge_cells("A5:F5")
            sheet.row_dimensions[5].height = 30
            sheet.append([])
            sheet.append(["Seat / Bench Number", "Block Number", "Class", "Department", "Subject", "Roll Number"])
            seats = {student["seat"]: (block, student) for block in room_blocks for student in block["students"]}
            for seat in range(1, room["capacity"] + 1):
                if seat in seats:
                    block, student = seats[seat]
                    exam = by_exam[student["timetable_id"]]
                    sheet.append([seat, block["block_number"], exam["class_name"], exam.get("department") or "Not specified", exam.get("subject", ""), student["roll"]])
                else:
                    sheet.append([seat, None, "EMPTY", None, None, None])
            _staff_style(sheet, 7, (22, 16, 24, 24, 48, 18))
    _staff_style(overview, 12, (29, 17, 15, 23, 24, 42, 58, 23, 17))
    workbook.active = 0


def write_seating_report(summaries, output_path=None, exam_info=None):
    """Export snapshots returned by generate_arrangements after its commit.

    Counts are student-exam attendances across slots, not unique people.
    Capacity used means occupied seats; unused includes wholly unused rooms.
    Room capacity and unused seats appear only on each room's first group row.
    """
    # Local import avoids a module cycle with the command-line entry point.
    from seating_allocator import compact_ranges

    summaries = list(summaries)
    target = Path(output_path) if output_path is not None else OUTPUT_DIR / "Seating_Arrangement.xlsx"
    eligible = sum(len(e["rolls"]) for _, exams, _, _ in summaries for e in exams)
    allocated = sum(len(a["rolls"]) for _, _, _, rows in summaries for a in rows)
    capacity = sum(r["capacity"] for _, _, rooms, _ in summaries for r in rooms)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Seating Arrangement"
    sheet.append(["EXAM SEATING ARRANGEMENT"])
    sheet.merge_cells("A1:J1")
    for label, value in [("Total Eligible Students", eligible),
                         ("Total Allocated Students", allocated),
                         ("Total Capacity Used", allocated),
                         ("Total Unused Seats", capacity - allocated)]:
        sheet.append([label, value])
    sheet.merge_cells("A6:J6")
    sheet["A6"] = ("Totals count student-exam attendances across slots. Capacity used = occupied seats. "
                   "Capacity and unused seats are shown once per room per slot, including empty rooms.")
    sheet.append([])
    sheet.append(["Exam Date", "Exam Time (Start)", "Exam Time (End)", "Room Number",
                  "Room Capacity", "Class", "Subject", "Allocated Roll Numbers",
                  "Number of Students Allocated", "Unused Seats"])
    for slot, exams, rooms, allocations in summaries:
        day, start, end = slot
        exam_by_id = {e["timetable_id"]: e for e in exams}
        by_room = defaultdict(list)
        for allocation in allocations:
            by_room[allocation["classroom_id"]].append(allocation)
        for room in rooms:
            groups = by_room[room["classroom_id"]]
            unused = room["capacity"] - sum(len(a["rolls"]) for a in groups)
            for index, allocation in enumerate(groups or [None]):
                exam = exam_by_id[allocation["timetable_id"]] if allocation else {}
                rolls = allocation["rolls"] if allocation else []
                sheet.append([_date(day), _time(start), _time(end), str(room["classroom_no"]),
                              room["capacity"] if index == 0 else None,
                              exam.get("class_name", ""), exam.get("subject", ""),
                              compact_ranges(rolls), len(rolls), unused if index == 0 else None])
                sheet.cell(sheet.max_row, 1).number_format = "dd-mmm-yyyy"
                for column in (2, 3):
                    sheet.cell(sheet.max_row, column).number_format = "hh:mm"

    navy = "203864"
    border = Border(*( [Side(style="thin", color="D5DCE4")] * 4))
    for row in sheet:
        for cell in row:
            cell.font = Font(name="Calibri", size=11, color="203040")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            # Imported identifiers and subjects must remain literal Excel text.
            if isinstance(cell.value, str):
                cell.data_type = "s"
            if cell.row >= 8 or 2 <= cell.row <= 5 and cell.column <= 2:
                cell.border = border
            if cell.row in (1, 8):
                cell.fill = PatternFill("solid", fgColor=navy)
                cell.font = Font(name="Calibri", size=14 if cell.row == 1 else 11,
                                 bold=True, color="FFFFFF")
            elif 2 <= cell.row <= 5 and cell.column <= 2:
                cell.font = Font(name="Calibri", size=11, bold=True)
                cell.fill = PatternFill("solid", fgColor="E8EFF7")
    for column, width in zip("ABCDEFGHIJ", (29, 19, 19, 17, 16, 24, 36, 60, 23, 16)):
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[6].height = 32
    sheet.row_dimensions[8].height = 34
    for row in range(9, sheet.max_row + 1):
        lines = max((len(str(sheet.cell(row, col).value or "")) + width - 1) // width
                    for col, width in ((6, 22), (7, 34), (8, 55)))
        sheet.row_dimensions[row].height = max(30, 15 * lines)
    sheet.freeze_panes = "F9"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = f"A8:J{sheet.max_row}"
    sheet.print_title_rows = "1:8"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True
    sheet.print_area = f"A1:J{sheet.max_row}"

    # Retain the original summary sheet for compatibility; staff-facing sheets open first.
    add_block_and_seat_sheets(workbook, summaries)
    if exam_info is not None:
        from college_report import add_college_report
        add_college_report(workbook, summaries, exam_info)

    # Apply to every output sheet after its layout is complete, preserving wrapping
    # and all other alignment settings along with the existing report design.
    from copy import copy
    for output_sheet in workbook.worksheets:
        for row in output_sheet:
            for cell in row:
                alignment = copy(cell.alignment)
                alignment.horizontal = "center"
                alignment.vertical = "center"
                cell.alignment = alignment

    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        # Replace only a complete workbook, preserving an older report on failure.
        with NamedTemporaryFile(dir=target.parent, suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        workbook.save(temporary)
        temporary.replace(target)
    finally:
        workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return target
