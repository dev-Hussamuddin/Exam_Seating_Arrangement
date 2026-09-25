"""Finalized workbooks, daily filtering, semester keys and college report checks."""

from datetime import date, datetime, time
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock, patch
import unittest

from openpyxl import load_workbook

import excel_importer as importer
import seating_allocator
from college_report import HEADERS, display_rolls
from create_input_templates import create_templates
from database import migrate_student_semester_schema
from seating_blocks import make_blocks
from seating_report import write_seating_report
from test_excel_importer import TestConnection
from three_file_input import (FILENAMES, STUDENT_HEADERS, ROOM_HEADERS, TIMETABLE_HEADERS,
                              read_three_file_input, three_file_input_available)
from three_file_workflow import generate_three_file_arrangements


class ThreeFileTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        create_templates(self.directory)
        self.connection = TestConnection()
        self.addCleanup(self.connection.db.close)
        self.connection.db.executescript("""
            ALTER TABLE student_batches ADD COLUMN semester TEXT;
            CREATE TABLE seating_arrangements (arrangement_id INTEGER PRIMARY KEY,
                timetable_id INTEGER, classroom_id INTEGER, roll_start INTEGER, roll_end INTEGER,
                roll_numbers TEXT NOT NULL, allocated_count INTEGER, block_number INTEGER, seat_numbers TEXT);
        """)

    def edit(self, file_index, change):
        path = self.directory / FILENAMES[file_index]
        book = load_workbook(path)
        try:
            change(book.active)
            book.save(path)
        finally:
            book.close()

    def import_data(self, data=None):
        data = data or read_three_file_input(self.directory)
        with patch.object(importer, "get_connection", return_value=self.connection), \
             patch.object(importer, "migrate_input_schema"), patch("database.migrate_student_semester_schema"):
            counts = importer.import_records(data["records"], semester_aware=True)
        return data, counts

    def test_exact_files_columns_example_counts_and_formatting(self):
        self.assertEqual({p.name for p in self.directory.iterdir()}, set(FILENAMES))
        for filename, header, row in zip(FILENAMES, (STUDENT_HEADERS, ROOM_HEADERS, TIMETABLE_HEADERS), (4, 4, 6)):
            book = load_workbook(self.directory / filename)
            try:
                self.assertEqual(tuple(c.value for c in book.active[row]), header)
                self.assertTrue(book.active.cell(row, 1).font.bold)
                self.assertEqual(book.active.page_setup.orientation, "landscape")
            finally:
                book.close()
        data = read_three_file_input(self.directory)
        self.assertEqual(len(data["records"]["Class Data"]), 6)
        self.assertEqual(len(data["exams"]), 4)
        self.assertEqual(sum(len(e["rolls"]) for e in data["exams"]), 225)
        self.assertEqual([r[1] for _, r in data["records"]["Classroom Data"]], [48, 40, 45, 55, 70])
        self.assertEqual(data["exam_info"]["start_time"], time(9))

    def test_ranges_lists_exclusions_and_leading_zero_display(self):
        self.edit(0, lambda s: (setattr(s["E5"], "value", "001-004, 008, 010-012"),
                                setattr(s["F5"], "value", "002,010")))
        data = read_three_file_input(self.directory)
        self.assertEqual(data["exams"][0]["rolls"], [1, 3, 4, 8, 11, 12])
        self.assertEqual(display_rolls(data["exams"][0]["rolls"], 3), "001, 003-004, 008, 011-012")

    def test_same_subject_disjoint_static_rows_combine(self):
        self.edit(0, lambda s: s.append(["Commerce", "F.Y.B.Com", "I", "Cost Accounting", "090,092-094", "093"]))
        data = read_three_file_input(self.directory)
        self.assertEqual(len(data["records"]["Class Data"]), 6)
        self.assertEqual(data["exams"][0]["rolls"][-3:], [90, 92, 94])

    def test_static_overlap_across_subjects_is_allowed_but_daily_overlap_is_not(self):
        read_three_file_input(self.directory)  # Business Law shares Cost Accounting's students, on a different day.
        self.edit(2, lambda s: s.append(["F.Y.B.Com", "I", "Business Law"]))
        with self.assertRaisesRegex(ValueError, "overlapping student"):
            read_three_file_input(self.directory)

    def test_invalid_rolls_exclusions_and_duplicate_static_group(self):
        for column, value in (("E5", "9-1"), ("E5", "1,,3"), ("F5", "999")):
            with self.subTest(value=value):
                book = load_workbook(self.directory / FILENAMES[0])
                old = book.active[column].value
                book.close()
                self.edit(0, lambda s: setattr(s[column], "value", value))
                with self.assertRaises(ValueError):
                    read_three_file_input(self.directory)
                self.edit(0, lambda s: setattr(s[column], "value", old))
        self.edit(0, lambda s: s.append(["Commerce", "F.Y.B.Com", "I", "Cost Accounting", "002-003", None]))
        with self.assertRaisesRegex(ValueError, "Duplicate/overlapping"):
            read_three_file_input(self.directory)

    def test_missing_triplet_member_never_falls_back_to_legacy(self):
        (self.directory / FILENAMES[1]).unlink()
        (self.directory / "Tables.xlsx").write_bytes(b"legacy placeholder")
        with patch.object(importer, "INPUT_DIR", self.directory), patch.object(importer, "get_connection") as connect:
            with self.assertRaisesRegex(ValueError, "Missing finalized"):
                importer.import_from_input()
            connect.assert_not_called()

    def test_default_discovery_prefers_triplet_alongside_legacy(self):
        (self.directory / "Tables.xlsx").write_bytes(b"must not be read")
        with patch.object(importer, "INPUT_DIR", self.directory), \
             patch.object(importer, "import_records", return_value={}) as run:
            importer.import_from_input()
        self.assertTrue(run.call_args.kwargs["semester_aware"])
        self.assertEqual(len(run.call_args.args[0]["Timetable"]), 4)

    def test_timetable_match_includes_semester_and_unknown_exam_fails(self):
        self.edit(0, lambda s: s.append(["Commerce", "F.Y.B.Com", "II", "Cost Accounting", "101-105", None]))
        self.edit(2, lambda s: setattr(s["B7"], "value", "II"))
        data, _ = self.import_data()
        self.assertEqual(data["exams"][0]["rolls"], [101, 102, 103, 104, 105])
        rows = self.connection.db.execute("SELECT semester, number_of_students FROM student_batches WHERE subject='Cost Accounting' ORDER BY semester").fetchall()
        self.assertEqual(rows, [("I", 78), ("II", 5)])
        summaries = generate_three_file_arrangements(self.connection, data)
        self.assertEqual(summaries[0][1][0]["semester"], "II")
        self.edit(2, lambda s: setattr(s["B7"], "value", "III"))
        with self.assertRaisesRegex(ValueError, "no matching student"):
            read_three_file_input(self.directory)

    def test_duplicate_daily_exam_rejected(self):
        self.edit(2, lambda s: s.append(["F.Y.B.Com", "I", "Cost Accounting"]))
        with self.assertRaisesRegex(ValueError, "duplicate exam"):
            read_three_file_input(self.directory)

    def test_capacity_checked_before_database(self):
        self.edit(1, lambda s: [setattr(s.cell(row, 2), "value", 1) for row in range(5, 10)])
        with patch.object(importer, "get_connection") as connect:
            with self.assertRaisesRegex(ValueError, "Insufficient capacity"):
                read_three_file_input(self.directory)
            connect.assert_not_called()

    def test_positive_capacity_duplicate_room_and_formula_validation(self):
        self.edit(1, lambda s: setattr(s["B5"], "value", 0))
        with self.assertRaises(ValueError):
            read_three_file_input(self.directory)
        self.edit(1, lambda s: (setattr(s["B5"], "value", 48), s.append(["201", 10])))
        with self.assertRaisesRegex(ValueError, "duplicate room"):
            read_three_file_input(self.directory)
        self.edit(1, lambda s: (s.delete_rows(10), setattr(s["B5"], "value", "=48")))
        with self.assertRaisesRegex(ValueError, "not formulas"):
            read_three_file_input(self.directory)

    def test_metadata_date_formats_optional_time_and_required_labels(self):
        self.edit(2, lambda s: (setattr(s["B1"], "value", "09/10/2026"), setattr(s["B4"], "value", None)))
        info = read_three_file_input(self.directory)["exam_info"]
        self.assertEqual(info["exam_date"], date(2026, 10, 9))
        self.assertIsNone(info["start_time"])
        self.edit(2, lambda s: setattr(s["A2"], "value", "Wrong label"))
        with self.assertRaisesRegex(ValueError, "top configuration"):
            read_three_file_input(self.directory)

    def test_repeat_import_and_allocation_preserve_ids_and_seats(self):
        data, first = self.import_data()
        self.assertEqual(first, dict(classes=4, batches=6, classrooms=5, timetable=4))
        generate_three_file_arrangements(self.connection, data)
        query = "SELECT timetable_id, classroom_id, roll_numbers, allocated_count, block_number, seat_numbers FROM seating_arrangements ORDER BY classroom_id, block_number, timetable_id"
        before = self.connection.db.execute(query).fetchall()
        _, second = self.import_data(data)
        self.assertEqual(second["timetable"], 0)
        generate_three_file_arrangements(self.connection, data)
        self.assertEqual(self.connection.db.execute(query).fetchall(), before)

    def test_only_current_timetable_and_rooms_participate(self):
        data, _ = self.import_data()
        db = self.connection.db
        db.execute("INSERT INTO classrooms VALUES (99, 'OLD', 999)")
        db.execute("INSERT INTO timetable VALUES (99, 1, 'Business Law', 'I', '2026-10-09', '09:00:00', '10:00:00', NULL)")
        db.execute("INSERT INTO seating_arrangements VALUES (99, 99, 99, 1, 1, '1', 1, NULL, NULL)")
        db.commit()
        old = db.execute("SELECT * FROM seating_arrangements WHERE arrangement_id=99").fetchone()
        summaries = generate_three_file_arrangements(self.connection, data)
        self.assertEqual(sum(len(a["rolls"]) for a in summaries[0][3]), 225)
        self.assertNotIn("Business Law", {e["subject"] for e in summaries[0][1]})
        self.assertNotIn("OLD", {r["classroom_no"] for r in summaries[0][2]})
        self.assertEqual(db.execute("SELECT * FROM seating_arrangements WHERE arrangement_id=99").fetchone(), old)

    def test_removed_daily_exam_is_not_allocated_on_next_run(self):
        data, _ = self.import_data()
        generate_three_file_arrangements(self.connection, data)
        self.edit(2, lambda s: s.delete_rows(7))
        updated, _ = self.import_data()
        summaries = generate_three_file_arrangements(self.connection, updated)
        self.assertEqual(len(summaries[0][1]), 3)
        self.assertEqual(sum(len(a["rolls"]) for a in summaries[0][3]), 147)
        self.assertNotIn("Cost Accounting", {e["subject"] for e in summaries[0][1]})

    def test_allocation_insert_failure_rolls_back_existing_rows(self):
        data, _ = self.import_data()
        generate_three_file_arrangements(self.connection, data)
        db = self.connection.db
        before = db.execute("SELECT * FROM seating_arrangements").fetchall()
        db.execute("CREATE TRIGGER fail_insert BEFORE INSERT ON seating_arrangements BEGIN SELECT RAISE(ABORT, 'test failure'); END")
        db.commit()
        with self.assertRaises(Exception):
            generate_three_file_arrangements(self.connection, data)
        self.assertEqual(db.execute("SELECT * FROM seating_arrangements").fetchall(), before)

    def test_import_failure_rolls_back_semester_batches_and_rooms(self):
        self.import_data()
        db = self.connection.db
        before = list(db.iterdump())
        data = read_three_file_input(self.directory)
        # Introduce a conflicting known department after other updates would have run.
        number, record = data["records"]["Class Data"][-1]
        data["records"]["Class Data"][-1] = (number, (record[0], "Wrong department", *record[2:]))
        with self.assertRaisesRegex(ValueError, "already belongs"):
            self.import_data(data)
        self.assertEqual(list(db.iterdump()), before)

    def test_college_layout_metadata_semesters_totals_and_global_blocks(self):
        data, _ = self.import_data()
        summaries = generate_three_file_arrangements(self.connection, data)
        path = write_seating_report(summaries, self.directory / "out" / "report.xlsx", data["exam_info"])
        book = load_workbook(path)
        try:
            sheet = book.active
            self.assertEqual(sheet.title, "Seating Arrangement")
            self.assertEqual(tuple(c.value for c in sheet[8]), HEADERS)
            self.assertEqual(sheet["A3"].value, data["exam_info"]["exam_name"])
            self.assertEqual(sheet["A4"].value, "October 2026")
            self.assertEqual(sheet["A6"].value, datetime(2026, 10, 9))
            self.assertEqual(sheet["A6"].number_format, "dd/mm/yyyy dddd")
            self.assertEqual(sheet["A7"].value, "Timing: 9:00 a.m. - 10:00 a.m.")
            self.assertEqual(sheet.cell(sheet.max_row, 7).value, 225)
            self.assertEqual(sum(sheet.cell(r, 7).value for r in range(9, sheet.max_row)), 225)
            self.assertEqual({sheet.cell(r, 4).value for r in range(9, sheet.max_row)}, {"I"})
            self.assertEqual([sheet.cell(r, 2).value for r in range(9, sheet.max_row) if sheet.cell(r, 2).value], list(range(1, 9)))
            self.assertEqual(book["Room Blocks"]["B9"].value, 8)
            self.assertEqual(sheet.max_column, 7)
            self.assertTrue(sheet["A8"].font.bold)
            self.assertEqual(sheet["F9"].border.bottom.style, "thin")
            self.assertEqual(sheet.page_setup.orientation, "landscape")
            self.assertIn("Seats 005", book.sheetnames)
            self.assertIn("Allocation Summary", book.sheetnames)
        finally:
            book.close()

    def test_63_students_in_one_room_through_three_file_workflow(self):
        self.edit(0, lambda s: (setattr(s["E5"], "value", "001-063"), setattr(s["F5"], "value", None)))
        self.edit(2, lambda s: s.delete_rows(8, 3))
        self.edit(1, lambda s: (setattr(s["B5"], "value", 70), s.delete_rows(6, 4)))
        data, _ = self.import_data()
        summaries = generate_three_file_arrangements(self.connection, data)
        blocks = make_blocks(summaries[0][3])
        self.assertEqual([len(block["students"]) for block in blocks], [30, 33])

    def test_template_creation_refuses_to_overwrite(self):
        before = [(self.directory / name).read_bytes() for name in FILENAMES]
        with self.assertRaises(FileExistsError):
            create_templates(self.directory)
        self.assertEqual([(self.directory / name).read_bytes() for name in FILENAMES], before)

    def test_semester_migration_only_adds_missing_column(self):
        connection = Mock()
        cursor = connection.cursor.return_value
        cursor.fetchone.return_value = None
        migrate_student_semester_schema(connection)
        cursor.execute.assert_any_call("ALTER TABLE student_batches ADD COLUMN semester VARCHAR(50) NULL")
        cursor.reset_mock()
        cursor.fetchone.return_value = ("exists",)
        migrate_student_semester_schema(connection)
        self.assertTrue(all(call.args[0].startswith("SHOW") for call in cursor.execute.call_args_list))

    def test_cli_mismatched_date_fails_before_database(self):
        with patch.object(seating_allocator, "get_connection") as connect:
            self.assertEqual(seating_allocator.main("2026-10-10", input_dir=self.directory), 1)
            connect.assert_not_called()


if __name__ == "__main__":
    unittest.main()
