"""One-day input, continuous blocks, real seat order and safe persistence/export."""

from copy import deepcopy
from datetime import date, datetime, time
from pathlib import Path
from tempfile import TemporaryDirectory
from io import BytesIO
import unittest
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

import excel_importer as importer
from database import migrate_block_schema
from one_day_input import STUDENT_HEADERS, ROOM_HEADERS
from seating_allocator import allocate_slot, generate_arrangements, validate_allocation
from seating_blocks import make_blocks
from seating_report import write_seating_report
from test_excel_importer import TestConnection
from test_seating_allocator import exam, rooms


class OneDayInputTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "OneDay.xlsx"
        self.book = Workbook()
        self.addCleanup(self.book.close)
        config = self.book.active
        config.title = "Exam Configuration"
        config.append(["Exam Date", date(2026, 10, 5)])
        config.append(["Exam Time", "12:00-01:00"])
        self.students = self.book.create_sheet("Student Subject Data")
        self.students.append(STUDENT_HEADERS)
        self.students.append(["SYCS", "CS", "Subject A", 1, 40, "3,8-9"])
        self.students.append(["SYCS", "CS", "Subject B", 41, 70, "45"])
        self.classrooms = self.book.create_sheet("Classroom Data")
        self.classrooms.append(ROOM_HEADERS)
        self.classrooms.append(["201", 48])
        self.classrooms.append(["202", 40])
        self.connection = TestConnection()
        self.addCleanup(self.connection.db.close)
        self.connection.db.execute("""CREATE TABLE seating_arrangements (
            arrangement_id INTEGER PRIMARY KEY, timetable_id INTEGER, classroom_id INTEGER,
            roll_start INTEGER, roll_end INTEGER, roll_numbers TEXT NOT NULL, allocated_count INTEGER,
            block_number INTEGER, seat_numbers TEXT)""")
        self.connection.db.commit()

    def read(self):
        self.book.save(self.path)
        return importer.read_workbook(self.path)

    def run_import(self):
        self.book.save(self.path)
        with patch.object(importer, "get_connection", return_value=self.connection), \
             patch.object(importer, "migrate_input_schema"):
            return importer.import_workbook(self.path)

    def test_single_date_subject_groups_and_calculated_counts(self):
        records = self.read()
        self.assertEqual([r[4] for _, r in records["Class Data"]], [37, 29])
        self.assertEqual({r[3] for _, r in records["Timetable"]}, {date(2026, 10, 5)})
        self.assertEqual([(r[4], r[5]) for _, r in records["Timetable"]], [(time(12), time(13))] * 2)
        self.assertEqual([r[1] for _, r in records["Classroom Data"]], [48, 40])

    def test_same_subject_disjoint_ranges_merge_without_filling_gaps(self):
        self.students.append(["SYCS", "CS", "Subject A", 90, 92, "91"])
        records = self.read()
        self.assertEqual(len(records["Class Data"]), 2)
        group = records["Class Data"][0][1]
        self.assertEqual(group[4], 39)
        self.assertNotIn("80", group[7].split(","))
        self.run_import()
        result = generate_arrangements(self.connection, one_day=True)
        self.assertEqual(sum(len(a["rolls"]) for a in result[0][3]), 68)

    def test_duplicate_and_overlapping_groups_rejected(self):
        for subject in ("Subject A", "Different subject"):
            with self.subTest(subject=subject):
                self.students.append(["SYCS", "CS", subject, 1, 4, None])
                with self.assertRaisesRegex(ValueError, "Overlapping"):
                    self.read()
                self.students.delete_rows(4)

    def test_exclusions_outside_range_and_missing_department_rejected(self):
        for cell, value in (("F2", "99"), ("B2", None), ("D2", 99)):
            with self.subTest(cell=cell):
                old = self.students[cell].value
                self.students[cell] = value
                with self.assertRaises(ValueError):
                    self.read()
                self.students[cell] = old

    def test_optional_time_round_trip_and_rerun(self):
        self.book["Exam Configuration"]["B2"] = None
        self.run_import()
        summaries = generate_arrangements(self.connection, one_day=True)
        self.assertEqual(summaries[0][0], ("2026-10-05", None, None))
        self.assertEqual(self.run_import()["timetable"], 0)
        generate_arrangements(self.connection, one_day=True)
        self.assertEqual(self.connection.db.execute("SELECT SUM(allocated_count) FROM seating_arrangements").fetchone()[0], 66)
        path = write_seating_report(summaries, Path(self.temp.name) / "report.xlsx")
        book = load_workbook(path)
        try:
            self.assertEqual(book.active["B3"].value, "Not configured")
        finally:
            book.close()

    def test_invalid_configuration_and_formula_rejected(self):
        for cell, value in (("B1", None), ("B1", "2026-02-30"), ("B2", "12:00"), ("B2", "=1+1")):
            with self.subTest(value=value):
                sheet = self.book["Exam Configuration"]
                old = sheet[cell].value
                sheet[cell] = value
                with self.assertRaises(ValueError):
                    self.read()
                sheet[cell] = old

    def test_department_fills_unknown_without_overwriting_known_department(self):
        db = self.connection.db
        db.execute("INSERT INTO classes VALUES (1, 'SYCS', '')")
        db.commit()
        self.run_import()
        self.assertEqual(db.execute("SELECT department FROM classes").fetchone()[0], "CS")
        self.students["B2"] = self.students["B3"] = "IT"
        with self.assertRaisesRegex(ValueError, "already belongs"):
            self.run_import()
        self.assertEqual(db.execute("SELECT department FROM classes").fetchone()[0], "CS")

    def test_selected_day_preserves_historical_allocations(self):
        self.run_import()
        db = self.connection.db
        db.execute("INSERT INTO timetable VALUES (99, 1, 'Subject A', NULL, '2026-10-04', '12:00:00', '13:00:00', NULL)")
        db.execute("INSERT INTO seating_arrangements VALUES (99, 99, 1, 1, 1, '1', 1, NULL, NULL)")
        db.commit()
        old = db.execute("SELECT * FROM seating_arrangements WHERE arrangement_id=99").fetchone()
        with self.assertRaisesRegex(ValueError, "Multiple exam dates"):
            generate_arrangements(self.connection, one_day=True)
        summaries = generate_arrangements(self.connection, exam_day="2026-10-05", one_day=True)
        self.assertEqual(len(summaries), 1)
        self.assertEqual(db.execute("SELECT * FROM seating_arrangements WHERE arrangement_id=99").fetchone(), old)

    def test_blocks_and_seats_persist_and_rerun_without_duplicates(self):
        self.run_import()
        generate_arrangements(self.connection, one_day=True)
        db = self.connection.db
        query = "SELECT classroom_id, block_number, allocated_count, roll_numbers, seat_numbers FROM seating_arrangements ORDER BY block_number"
        first = db.execute(query).fetchall()
        self.assertEqual([r[1] for r in first], [1, 2, 2, 3])
        self.assertEqual(db.execute("SELECT block_number, SUM(allocated_count) FROM seating_arrangements GROUP BY block_number ORDER BY block_number").fetchall(),
                         [(1, 30), (2, 18), (3, 18)])
        self.assertEqual(sum(r[2] for r in first), 66)
        for room_id in (1, 2):
            seats = [int(s) for r in first if r[0] == room_id for s in r[4].split(",")]
            self.assertEqual(sorted(seats), list(range(1, len(seats) + 1)))
        generate_arrangements(self.connection, one_day=True)
        self.assertEqual(db.execute(query).fetchall(), first)

    def test_mixed_subjects_share_persisted_block_and_remainder(self):
        self.classrooms["B2"] = 80
        self.run_import()
        generate_arrangements(self.connection, one_day=True)
        db = self.connection.db
        self.assertEqual(db.execute("SELECT block_number, SUM(allocated_count) FROM seating_arrangements GROUP BY block_number ORDER BY block_number").fetchall(),
                         [(1, 30), (2, 36)])
        self.assertEqual(db.execute("SELECT COUNT(DISTINCT timetable_id) FROM seating_arrangements WHERE block_number=2").fetchone()[0], 2)
        self.assertEqual(db.execute("SELECT SUM(allocated_count) FROM seating_arrangements").fetchone()[0], 66)

    def test_capacity_failure_preserves_previous_blocks_and_seats(self):
        self.run_import()
        generate_arrangements(self.connection, one_day=True)
        db = self.connection.db
        before = db.execute("SELECT * FROM seating_arrangements").fetchall()
        db.execute("UPDATE classrooms SET capacity=1")
        db.commit()
        with self.assertRaisesRegex(ValueError, "Insufficient capacity"):
            generate_arrangements(self.connection, one_day=True)
        self.assertEqual(db.execute("SELECT * FROM seating_arrangements").fetchall(), before)

    def test_untimed_and_timed_slots_cannot_silently_share_rooms(self):
        self.run_import()
        db = self.connection.db
        db.execute("INSERT INTO timetable VALUES (99, 1, 'Subject A', NULL, '2026-10-05', NULL, NULL, NULL)")
        db.commit()
        with self.assertRaisesRegex(ValueError, "untimed exam"):
            generate_arrangements(self.connection, one_day=True)


class BlockAndSeatTests(unittest.TestCase):
    def test_legacy_college_disjoint_subject_rows_and_overlap(self):
        book = Workbook()
        try:
            sheet = book.active
            for name, column in (("Class Data", 1), ("Classroom Data", 8), ("Timetable", 11)):
                sheet.cell(1, column, name)
                for col, header in enumerate(importer.COLLEGE_HEADERS[name], column):
                    sheet.cell(2, col, header)
            for row, values in enumerate([["SYCS", "Maths", "1-4", "2", 3],
                                          ["SYCS", "Maths", "10-12", "11", 2]], 3):
                for col, value in enumerate(values, 1):
                    sheet.cell(row, col, value)
            sheet.cell(3, 8, "201")
            sheet.cell(3, 9, 48)
            for col, value in enumerate(["SYCS", "Maths", "09:00-10:00", date(2026, 10, 5)], 11):
                sheet.cell(3, col, value)
            with BytesIO() as stream:
                book.save(stream)
                stream.seek(0)
                records = importer.read_workbook(stream)
            self.assertEqual(len(records["Class Data"]), 1)
            group = records["Class Data"][0][1]
            self.assertEqual(group[4:6], (5, "2,11"))
            self.assertEqual(group[7], "1,2,3,4,10,11,12")
            sheet.cell(4, 3, "4-12")
            with BytesIO() as stream:
                book.save(stream)
                stream.seek(0)
                with self.assertRaisesRegex(ValueError, "overlapping"):
                    importer.read_workbook(stream)
        finally:
            book.close()

    def test_block_boundary_and_continuous_room_numbering(self):
        allocations = allocate_slot([exam(1, "CS", 85)], rooms(48, 40))
        blocks = make_blocks(allocations)
        self.assertEqual([b["block_number"] for b in blocks], [1, 2, 3])
        self.assertEqual([len(b["students"]) for b in blocks], [30, 18, 37])
        self.assertEqual([b["classroom_id"] for b in blocks], [1, 1, 2])

    def test_63_students_in_one_room_make_exactly_30_and_33(self):
        allocations = allocate_slot([exam(1, "CS", 63)], rooms(70))
        before = deepcopy(allocations)
        blocks = make_blocks(allocations)
        self.assertEqual(len(blocks), 2)
        self.assertEqual([len(b["students"]) for b in blocks], [30, 33])
        self.assertEqual([[s["seat"] for s in b["students"]] for b in blocks],
                         [list(range(1, 31)), list(range(31, 64))])
        self.assertEqual(allocations, before)

    def test_target_rounding_absorbs_small_remainders_and_empty_rooms(self):
        for total, expected in ((1, [1]), (30, [30]), (31, [31]), (44, [44]),
                                (45, [30, 15]), (60, [30, 30]), (93, [30, 30, 33])):
            with self.subTest(total=total):
                blocks = make_blocks(allocate_slot([exam(1, "CS", total)], rooms(100, 20)), first_block=7)
                self.assertEqual([len(b["students"]) for b in blocks], expected)
                self.assertEqual([b["block_number"] for b in blocks], list(range(7, 7 + len(expected))))
        self.assertEqual(make_blocks([]), [])

    def test_sequential_group_order_is_preserved_in_seat_map(self):
        allocations = allocate_slot([exam(1, "CS", 4), exam(2, "Science", 4)], rooms(8))
        self.assertEqual(allocations[0]["seats"], [1, 2, 3, 4])
        self.assertEqual(allocations[1]["seats"], [5, 6, 7, 8])
        blocks = make_blocks(allocations)
        self.assertEqual(len(blocks), 1)
        self.assertEqual([s["seat"] for s in blocks[0]["students"]], list(range(1, 9)))
        self.assertEqual([s["class_id"] for s in blocks[0]["students"]], [1] * 4 + [2] * 4)
        self.assertEqual([g["seats"] for g in blocks[0]["groups"]], [[1, 2, 3, 4], [5, 6, 7, 8]])

    def test_99_cs_then_19_ba_across_rooms_and_college_report(self):
        exams = [{**exam(1, "CS", 99), "subject": "Computing", "semester": "I"},
                 {**exam(2, "Arts", 19), "subject": "History", "semester": "I"}]
        classroom = rooms(60, 60)
        classroom[0]["classroom_no"] = "201"
        classroom[1]["classroom_no"] = "202"
        allocations = allocate_slot(exams, classroom)
        self.assertEqual([(a["classroom_id"], a["class_id"], len(a["rolls"])) for a in allocations],
                         [(1, 1, 60), (2, 1, 39), (2, 2, 19)])
        self.assertEqual([a["rolls"] for a in allocations],
                         [list(range(1, 61)), list(range(61, 100)), list(range(1, 20))])
        self.assertEqual([a["seats"] for a in allocations],
                         [list(range(1, 61)), list(range(1, 40)), list(range(40, 59))])
        blocks = make_blocks(allocations)
        self.assertEqual([b["block_number"] for b in blocks], [1, 2, 3, 4])
        self.assertEqual([len(b["students"]) for b in blocks], [30, 30, 30, 28])
        self.assertEqual([[len(g["rolls"]) for g in b["groups"]] for b in blocks],
                         [[30], [30], [30], [9, 19]])
        info = dict(exam_name="Example", month_year="October 2026", exam_date=date(2026, 10, 5), timing="09:00-10:00")
        summaries = [((date(2026, 10, 5), time(9), time(10)), exams, classroom, allocations)]
        with TemporaryDirectory() as directory:
            book = load_workbook(write_seating_report(summaries, Path(directory) / "report.xlsx", exam_info=info))
            try:
                rows = list(book.active.iter_rows(min_row=9, max_row=13, values_only=True))
                self.assertEqual([(r[2], r[4], r[5], r[6]) for r in rows],
                                 [("Class1", "Computing", "1-30", 30),
                                  ("Class1", "Computing", "31-60", 30),
                                  ("Class1", "Computing", "61-90", 30),
                                  ("Class1", "Computing", "91-99", 9),
                                  ("Class2", "History", "1-19", 19)])
                self.assertEqual(book.active.cell(14, 7).value, 118)
            finally:
                book.close()

    def test_same_class_distinct_subject_groups_allowed(self):
        first = exam(1, "CS", 4)
        second = {**exam(1, "CS", 10, start=7), "timetable_id": 2}
        self.assertEqual(sum(len(a["rolls"]) for a in allocate_slot([first, second], rooms(8))), 8)
        second["rolls"] = [4, 7]
        with self.assertRaisesRegex(ValueError, "Overlapping"):
            allocate_slot([first, second], rooms(8))

    def test_excluded_rolls_absent_from_every_block(self):
        allocations = allocate_slot([exam(1, "CS", 65, "3,30,31,61")], rooms(40, 30))
        blocks = make_blocks(allocations)
        rolls = [s["roll"] for b in blocks for s in b["students"]]
        self.assertEqual(len(rolls), 61)
        self.assertEqual(len(set(rolls)), 61)
        self.assertFalse({3, 30, 31, 61} & set(rolls))

    def test_duplicate_or_missing_seats_rejected(self):
        exams, classroom = [exam(1, "CS", 4)], rooms(4)
        for invalid in ([1, 1, 3, 4], [1, 2, 3], [0, 1, 2, 3], [1, 2, 3, 5]):
            allocations = allocate_slot(exams, classroom)
            allocations[0]["seats"] = invalid
            with self.assertRaises(ValueError):
                validate_allocation(exams, classroom, allocations)

    def test_block_migration_is_additive_and_repeatable(self):
        connection = Mock()
        cursor = connection.cursor.return_value
        cursor.fetchone.return_value = None
        migrate_block_schema(connection)
        sql = [call.args[0] for call in cursor.execute.call_args_list]
        self.assertEqual(sum(s.startswith("ALTER TABLE") for s in sql), 2)
        self.assertFalse(any("DELETE" in s or "UPDATE" in s or "DROP" in s for s in sql))
        cursor.reset_mock()
        cursor.fetchone.return_value = ("exists",)
        migrate_block_schema(connection)
        self.assertTrue(all(call.args[0].startswith("SHOW") for call in cursor.execute.call_args_list))

    def test_report_blocks_seats_totals_and_print_settings(self):
        exams = [{**exam(1, "CS", 65, "3,30"), "subject": "=Literal Subject"}]
        classroom = rooms(48, 40, 5)
        allocations = allocate_slot(exams, classroom)
        summaries = [((date(2026, 10, 5), time(12), time(13)), exams, classroom, allocations)]
        before = deepcopy(summaries)
        with TemporaryDirectory() as directory:
            path = write_seating_report(summaries, Path(directory) / "nested" / "report.xlsx")
            book = load_workbook(path)
            try:
                self.assertEqual(book.active.title, "Room Blocks")
                sheet = book.active
                self.assertEqual(sheet["B2"].value, datetime(2026, 10, 5))
                self.assertEqual(sheet["B3"].value, time(12))
                self.assertEqual([sheet.cell(r, 2).value for r in range(5, 10)], [63, 63, 63, 30, 3])
                self.assertEqual([sheet.cell(r, 3).value for r in range(13, 16)], [1, 2, 3])
                self.assertEqual(sheet["E13"].value, "CS")
                self.assertEqual(sheet["F13"].data_type, "s")
                self.assertTrue(sheet["A12"].font.bold)
                self.assertEqual(sheet["G13"].border.bottom.style, "thin")
                self.assertEqual(sheet.page_setup.orientation, "landscape")
                self.assertEqual(sheet.print_title_rows, "$1:$12")
                first, second, empty = (book[f"Seats {n:03d}"] for n in (1, 2, 3))
                self.assertEqual(first.max_row, 55)
                self.assertEqual(first["F10"].value, 4)  # Seat 3 skips excluded roll 3.
                self.assertEqual(second["B8"].value, 3)
                self.assertEqual(second["C23"].value, "EMPTY")
                self.assertEqual(empty["F4"].value, 5)
                self.assertTrue(all(empty.cell(r, 3).value == "EMPTY" for r in range(8, 13)))
            finally:
                book.close()
        self.assertEqual(summaries, before)

    def test_report_numbers_continue_across_slots(self):
        exams, classroom = [exam(1, "CS", 4)], rooms(4)
        allocations = allocate_slot(exams, classroom)
        summaries = [((date(2026, 10, 5), time(h), time(h + 1)), exams, classroom, allocations) for h in (9, 12)]
        with TemporaryDirectory() as directory:
            path = write_seating_report(summaries, Path(directory) / "report.xlsx")
            book = load_workbook(path)
            try:
                self.assertEqual(book["Seats 001"]["B8"].value, 1)
                self.assertEqual(book["Seats 002"]["B8"].value, 2)
            finally:
                book.close()

    def test_63_mixed_students_report_two_blocks_with_correct_seat_identity(self):
        exams = [{**exam(1, "CS", 32), "subject": "Computing"},
                 {**exam(2, "Science", 31), "subject": "Physics"}]
        classroom = rooms(70)
        allocations = allocate_slot(exams, classroom)
        summaries = [((date(2026, 10, 5), time(9), time(10)), exams, classroom, allocations)]
        with TemporaryDirectory() as directory:
            path = write_seating_report(summaries, Path(directory) / "report.xlsx")
            book = load_workbook(path)
            try:
                self.assertEqual(book.active["B9"].value, 2)
                rows = list(book.active.iter_rows(min_row=13, values_only=True))
                self.assertEqual([row[2] for row in rows], [1, 2, 2])
                self.assertEqual([sum(row[7] for row in rows if row[2] == n) for n in (1, 2)], [30, 33])
                seat_rows = list(book["Seats 001"].iter_rows(min_row=8, max_row=70, values_only=True))
                self.assertEqual([row[1] for row in seat_rows], [1] * 30 + [2] * 33)
                self.assertEqual(len({(row[2], row[5]) for row in seat_rows}), 63)
                self.assertEqual([row[2] for row in seat_rows[:4]], ["Class1"] * 4)
                self.assertEqual([row[5] for row in seat_rows[:4]], [1, 2, 3, 4])
                self.assertEqual(book["Seats 001"]["C71"].value, "EMPTY")
            finally:
                book.close()


if __name__ == "__main__":
    unittest.main()
