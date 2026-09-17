"""Offline importer checks using disposable workbooks and an in-memory database."""

import io
import sqlite3
import unittest
from contextlib import redirect_stdout
from datetime import date, datetime, time
from tempfile import TemporaryDirectory
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

import excel_importer as importer


class TestConnection:
    """Adapt only parameter syntax/types; exercise real SQL and transactions.

    These tests do not replace a live MySQL integration check.
    """

    def __init__(self):
        self.db = sqlite3.connect(":memory:")
        self.db.executescript("""
            CREATE TABLE classes (class_id INTEGER PRIMARY KEY, class_name TEXT UNIQUE, department TEXT);
            CREATE TABLE student_batches (batch_id INTEGER PRIMARY KEY, class_id INTEGER,
                academic_year TEXT, roll_start INTEGER, roll_end INTEGER,
                number_of_students INTEGER, non_included_rolls TEXT, subject TEXT, roll_numbers TEXT);
            CREATE TABLE classrooms (classroom_id INTEGER PRIMARY KEY, classroom_no TEXT UNIQUE, capacity INTEGER);
            CREATE TABLE timetable (timetable_id INTEGER PRIMARY KEY, class_id INTEGER,
                subject TEXT, semester TEXT, exam_date TEXT, start_time TEXT, end_time TEXT, time_text TEXT);
        """)

    def start_transaction(self):
        self.db.execute("BEGIN")

    def cursor(self, buffered=True):
        return TestCursor(self.db.cursor())

    def commit(self):
        self.db.commit()

    def rollback(self):
        self.db.rollback()

    def close(self):
        pass  # Keep the test database open for assertions.


class TestCursor:
    def __init__(self, cursor):
        self.cursor = cursor

    def execute(self, sql, values=()):
        values = tuple(v.isoformat() if isinstance(v, (date, time)) else v for v in values)
        self.cursor.execute(sql.replace("%s", "?"), values)

    def __getattr__(self, name):
        return getattr(self.cursor, name)


class ImporterTests(unittest.TestCase):
    def setUp(self):
        self.connection = TestConnection()
        self.addCleanup(self.connection.db.close)
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "exam_input.xlsx"
        self.workbook = Workbook()
        self.workbook.remove(self.workbook.active)
        self.addCleanup(self.workbook.close)
        for name, headers in importer.HEADERS.items():
            self.workbook.create_sheet(name).append(headers)
        self.workbook["Class Data"].append(["FYCS", "CS", "1-60", 60, "5,12"])
        self.workbook["Class Data"].append(["SYCS", "CS", "1-55", 55, 10])
        self.workbook["Classroom Data"].append(["101", 30])
        self.workbook["Timetable"].append(["FYCS", "Maths", None, datetime(2026, 10, 1), time(9), time(11)])
        self.workbook["Timetable"].append(["SYCS", "Science", "3", "2026-10-01", "09:00", "11:00"])

    def save_single_sheet(self, offset=0, reverse=False):
        """Pack editable test tables onto one sheet with independently shifted rows."""
        combined = Workbook()
        sheet = combined.active
        sheet.title = "College Input"
        col = 1 + offset
        sources = list(self.workbook.worksheets)
        if reverse:
            sources.reverse()
        for index, source in enumerate(sources):
            title_row = 2 + offset + index
            sheet.cell(title_row, col, source.title)
            sheet.merge_cells(start_row=title_row, start_column=col,
                              end_row=title_row, end_column=col + source.max_column - 1)
            for r, row in enumerate(source.iter_rows(values_only=True), title_row + 2):
                for c, value in enumerate(row, col):
                    sheet.cell(r, c, value)
            col += source.max_column + 2
        combined.save(self.path)
        combined.close()

    def run_import(self):
        self.save_single_sheet()
        with patch.object(importer, "get_connection", return_value=self.connection), patch.object(importer, "migrate_input_schema"):
            return importer.import_workbook(self.path)

    def test_repeat_import_updates_without_duplicates(self):
        first = self.run_import()
        self.assertEqual(first, dict(classes=2, batches=2, classrooms=1, timetable=2))
        self.workbook["Class Data"]["D2"] = 59
        self.workbook["Classroom Data"]["B2"] = 40
        second = self.run_import()
        self.assertEqual(second["classes"], 0)
        self.assertEqual(second["timetable"], 0)
        db = self.connection.db
        self.assertEqual(db.execute("SELECT COUNT(*) FROM student_batches").fetchone()[0], 2)
        self.assertEqual(db.execute("SELECT capacity FROM classrooms").fetchone()[0], 40)
        self.assertEqual(db.execute("SELECT academic_year, number_of_students, non_included_rolls FROM student_batches WHERE class_id=1").fetchone(),
                         (importer.ACADEMIC_YEAR, 59, "5,12"))
        self.assertEqual(db.execute("SELECT c.class_name, t.subject FROM timetable t JOIN classes c ON c.class_id=t.class_id ORDER BY t.timetable_id").fetchall(),
                         [("FYCS", "Maths"), ("SYCS", "Science")])

    def test_late_failure_rolls_back_inserts_and_updates(self):
        self.run_import()
        self.connection.db.execute("CREATE TRIGGER fail_new_exam BEFORE INSERT ON timetable BEGIN SELECT RAISE(ABORT, 'test insert failure'); END")
        self.connection.db.commit()
        before = list(self.connection.db.iterdump())
        self.workbook["Classroom Data"]["B2"] = 99
        self.workbook["Class Data"].append(["TYCS", "CS", "1-40", 40, None])
        self.workbook["Timetable"].append(["UNKNOWN", "Maths", None, "2026-10-01", "09:00", "11:00"])
        with self.assertRaisesRegex(sqlite3.IntegrityError, "test insert failure"):
            self.run_import()
        self.assertEqual(list(self.connection.db.iterdump()), before)

    def test_invalid_inputs_rejected(self):
        cases = [("Class Data", "C2", "60-1"), ("Classroom Data", "B2", 0),
                 ("Class Data", "D2", 1.5), ("Class Data", "E2", "61"),
                 ("Timetable", "B2", None), ("Timetable", "D2", "bad date"),
                 ("Timetable", "E2", "25:00"), ("Timetable", "F2", "08:00"),
                 ("Class Data", "D2", "=60"), ("Class Data", "A1", "Wrong header")]
        for sheet, cell, value in cases:
            with self.subTest(sheet=sheet, cell=cell, value=value):
                original = self.workbook[sheet][cell].value
                self.workbook[sheet][cell] = value
                with self.assertRaises(ValueError):
                    self.run_import()
                self.workbook[sheet][cell] = original
        self.assertEqual(self.connection.db.execute("SELECT COUNT(*) FROM classes").fetchone()[0], 0)

    def test_master_department_conflict(self):
        self.run_import()
        self.workbook["Class Data"]["B2"] = "IT"
        with self.assertRaisesRegex(ValueError, "already belongs"):
            self.run_import()

    def test_missing_table(self):
        del self.workbook["Timetable"]
        with self.assertRaisesRegex(ValueError, "three table headings"):
            self.run_import()

    def test_shifted_and_reordered_tables(self):
        self.save_single_sheet()
        baseline = importer.read_workbook(self.path)
        self.save_single_sheet(offset=4, reverse=True)
        shifted = importer.read_workbook(self.path)
        self.assertEqual({name: [record for _, record in rows] for name, rows in baseline.items()},
                         {name: [record for _, record in rows] for name, rows in shifted.items()})

    def test_blank_rows_and_different_table_lengths(self):
        self.workbook["Class Data"].insert_rows(3)
        self.workbook["Timetable"].append(["FYCS", "English", None, "2026-10-02", "09:00", "11:00"])
        counts = self.run_import()
        self.assertEqual(counts, dict(classes=2, batches=2, classrooms=1, timetable=3))

    def test_reordered_headers_and_whitespace(self):
        sheet = self.workbook["Class Data"]
        for row in sheet:
            row[0].value, row[1].value = row[1].value, row[0].value
        sheet["C1"] = "  Roll Nos.\n(Range)  "
        self.assertEqual(self.run_import()["classes"], 2)

    def test_duplicate_title_rejected_before_database(self):
        self.workbook["Class Data"].append(["Class Data"])
        with patch.object(importer, "get_connection") as connect:
            with self.assertRaisesRegex(ValueError, "heading must appear exactly once"):
                self.run_import()
            connect.assert_not_called()

    def test_duplicate_header_rejected(self):
        self.workbook["Class Data"]["F1"] = "Class"
        with self.assertRaisesRegex(ValueError, "duplicate column header"):
            self.run_import()

    def test_file_selection(self):
        with patch.object(importer, "INPUT_DIR", Path(self.temp.name)):
            output = io.StringIO()
            with redirect_stdout(output):
                importer.import_from_input()
            self.assertIn("No .xlsx workbook found", output.getvalue())
            self.workbook.save(self.path)
            (Path(self.temp.name) / "~$exam_input.xlsx").touch()
            with patch.object(importer, "import_workbook", return_value=dict(classes=0, batches=0, classrooms=0, timetable=0)) as run:
                with redirect_stdout(io.StringIO()):
                    importer.import_from_input()
                run.assert_called_once_with(self.path)
            (Path(self.temp.name) / "second.xlsx").touch()
            with self.assertRaisesRegex(ValueError, "exactly one"):
                importer.import_from_input()


if __name__ == "__main__":
    unittest.main()
