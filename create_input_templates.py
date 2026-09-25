"""Create the finalized example workbooks without overwriting any existing input."""

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import INPUT_DIR
from three_file_input import FILENAMES, STUDENT_HEADERS, ROOM_HEADERS, TIMETABLE_HEADERS


EXAMPLE_STUDENTS = (
    ("Commerce", "F.Y.B.Com", "I", "Cost Accounting", "001-080", "015, 023"),
    ("Commerce", "F.Y.B.Com", "I", "Business Law", "001-080", "—"),
    ("Computer Science", "F.Y.B.Sc (Computer Science)", "I", "Database Management", "001-060", "009, 024"),
    ("Computer Science", "F.Y.B.Sc (Computer Science)", "I", "Mathematics", "001-060", "—"),
    ("Biotechnology", "F.Y.B.Sc (Biotechnology)", "I", "Physics", "001-020, 021-040", "012"),
    ("Arts", "F.Y.B.A", "I", "History", "001-050", "—"),
)


def _style(sheet, header, widths, color):
    side = Side(style="thin", color="BCC9D4")
    for row in sheet:
        for cell in row:
            cell.font = Font(name="Calibri", size=12, color="203040")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if isinstance(cell.value, str):
                cell.data_type = "s"
            if cell.row >= header:
                cell.border = Border(left=side, right=side, top=side, bottom=side)
                if cell.row > header and (cell.row - header) % 2:
                    cell.fill = PatternFill("solid", fgColor="F2F6FA")
            if cell.row == header:
                cell.fill = PatternFill("solid", fgColor=color)
                cell.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[row[0].row].height = 44 if row[0].row >= header else 26
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = f"A{header + 1}"
    last = get_column_letter(len(widths))
    sheet.auto_filter.ref = f"A{header}:{last}{sheet.max_row}"
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = f"1:{header}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_area = f"A1:{last}{sheet.max_row}"


def create_templates(directory=INPUT_DIR):
    directory = Path(directory)
    paths = [directory / name for name in FILENAMES]
    if any(path.exists() for path in paths):
        raise FileExistsError("Existing finalized inputs are preserved; use an empty destination to create examples.")
    directory.mkdir(parents=True, exist_ok=True)
    legacy = directory / "Tables.xlsx"
    if legacy.exists():
        from excel_importer import read_workbook
        room_rows = [record for _, record in read_workbook(legacy)["Classroom Data"]]
    else:
        room_rows = [("201", 48), ("202", 40), ("203", 45), ("204", 55), ("205", 70)]
    for index, (name, headers, rows, widths, color) in enumerate((
        ("Students & Subject Data", STUDENT_HEADERS, EXAMPLE_STUDENTS, (24, 35, 14, 32, 32, 23), "124D79"),
        ("Classroom Data", ROOM_HEADERS, room_rows, (24, 45), "176245"),
    )):
        book = Workbook()
        try:
            sheet = book.active
            sheet.title = name
            sheet.append([name + " — Static"])
            sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
            sheet.append(["Example student data — replace with actual college records." if index == 0 else
                          "One student per bench/seat. Existing room capacities are retained."])
            sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
            sheet.append([])
            sheet.append(headers)
            for row in rows:
                sheet.append(row)
            _style(sheet, 4, widths, color)
            sheet["A1"].font = Font(name="Calibri", size=18, bold=True, color=color)
            sheet.row_dimensions[1].height = 32
            if index == 1:
                sheet.row_dimensions[4].height = 56
            book.save(paths[index])
        finally:
            book.close()
    book = Workbook()
    try:
        sheet = book.active
        sheet.title = "Timetable"
        for row in (("Exam Date", date(2026, 10, 9)), ("Exam Name", "F.Y. - Regular/Additional Examination"),
                    ("Month / Year", "October 2026"), ("Timing", "9:00 a.m. - 10:00 a.m.")):
            sheet.append(row)
            sheet.merge_cells(start_row=sheet.max_row, start_column=2, end_row=sheet.max_row, end_column=3)
        sheet.append(["Example exam day — edit the details above and the participating exams below."])
        sheet.merge_cells("A5:C5")
        sheet.append(TIMETABLE_HEADERS)
        for index in (0, 2, 4, 5):
            row = EXAMPLE_STUDENTS[index]
            sheet.append([row[1], row[2], row[3]])
        _style(sheet, 6, (40, 20, 42), "93721B")
        for number in range(1, 5):
            sheet.cell(number, 1).font = Font(name="Calibri", size=12, bold=True)
            for cell in sheet[number]:
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
        sheet["B1"].number_format = "dd/mm/yyyy"
        sheet.row_dimensions[5].height = 34
        book.save(paths[2])
    finally:
        book.close()
    return paths


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--directory", type=Path, default=INPUT_DIR)
    for path in create_templates(parser.parse_args().directory):
        print(path)
