"""The finalized seven-column college seating sheet (no teacher fields)."""

from itertools import groupby
from math import ceil
import textwrap

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import COLLEGE_NAME, COLLEGE_SUBTITLE
from seating_blocks import make_blocks


HEADERS = ("Room No.", "Block No.", "Class", "Semester", "Subject", "Roll Nos.", "Grand Total")


def display_rolls(rolls, width=0):
    """Preserve numeric leading-zero presentation without changing student IDs."""
    parts = []
    for _, consecutive in groupby(enumerate(rolls), lambda pair: pair[1] - pair[0]):
        values = [pair[1] for pair in consecutive]
        first, last = str(values[0]).zfill(width), str(values[-1]).zfill(width)
        parts.append(first if len(values) == 1 else f"{first}-{last}")
    return ", ".join(parts)


def add_college_report(workbook, summaries, info):
    # Keep the old summary, but the finalized college layout is the main sheet.
    workbook["Seating Arrangement"].title = "Allocation Summary"
    sheet = workbook.create_sheet("Seating Arrangement", 0)
    for number, value in enumerate((COLLEGE_NAME, COLLEGE_SUBTITLE, info["exam_name"],
                                     info["month_year"], "Seating Arrangement", info["exam_date"],
                                     "Timing: " + (info["timing"] or "Not configured")), 1):
        sheet.append([value])
        sheet.merge_cells(start_row=number, start_column=1, end_row=number, end_column=7)
        cell = sheet.cell(number, 1)
        cell.font = Font(name="Times New Roman", size=18 if number == 1 else 13, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if isinstance(value, str):
            cell.data_type = "s"
        sheet.row_dimensions[number].height = 30 if number in (1, 3) else 22
    sheet["A6"].number_format = "dd/mm/yyyy dddd"
    sheet.append(HEADERS)
    widths = (12, 12, 34, 18, 34, 48, 14)
    edge = Side(style="thin", color="8A8A8A")
    for cell in sheet[8]:
        cell.font = Font(name="Times New Roman", size=12, bold=True)
        cell.fill = PatternFill("solid", fgColor="F4CCD9")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
    sheet.row_dimensions[8].height = 28
    next_block, total = 1, 0
    for _, exams, rooms, allocations in summaries:
        blocks = make_blocks(allocations, next_block)
        next_block += len(blocks)
        by_exam = {exam["timetable_id"]: exam for exam in exams}
        for room in rooms:
            room_start = sheet.max_row + 1
            room_blocks = [block for block in blocks if block["classroom_id"] == room["classroom_id"]]
            if not room_blocks:
                sheet.append([str(room["classroom_no"]), None, "Unused", None, None, None, 0])
            for block in room_blocks:
                block_start = sheet.max_row + 1
                for group in block["groups"]:
                    exam = by_exam[group["timetable_id"]]
                    count = len(group["rolls"])
                    total += count
                    sheet.append([str(room["classroom_no"]), block["block_number"], exam["class_name"],
                                  exam.get("semester", ""), exam.get("subject", ""),
                                  display_rolls(group["rolls"], exam.get("roll_width", 0)), count])
                if sheet.max_row > block_start:
                    sheet.merge_cells(start_row=block_start, start_column=2, end_row=sheet.max_row, end_column=2)
            if sheet.max_row > room_start:
                sheet.merge_cells(start_row=room_start, start_column=1, end_row=sheet.max_row, end_column=1)
    data_end = sheet.max_row
    sheet.append(["Grand Total", None, None, None, None, None, total])
    sheet.merge_cells(start_row=sheet.max_row, start_column=1, end_row=sheet.max_row, end_column=6)
    for row in sheet.iter_rows(min_row=9):
        for cell in row:
            cell.font = Font(name="Times New Roman", size=12, bold=cell.column <= 2 or cell.row > data_end)
            cell.alignment = Alignment(horizontal="center" if cell.column in (1, 2, 4, 5, 6, 7) else "left",
                                       vertical="center", wrap_text=True)
            cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
            if cell.column == 1 or cell.row > data_end:
                cell.fill = PatternFill("solid", fgColor="FBEAF0")
            if isinstance(cell.value, str):
                cell.data_type = "s"
        lines = max(max(1, len(textwrap.wrap(str(cell.value or ""), width=max(5, widths[cell.column - 1] - 5)))) for cell in row)
        sheet.row_dimensions[row[0].row].height = max(34, ceil(lines * 16 + 8))
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "C9"
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = "1:8"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_area = f"A1:G{sheet.max_row}"
    sheet.print_options.horizontalCentered = True
    sheet.oddFooter.center.text = "Page &P of &N"
    workbook.active = 0
