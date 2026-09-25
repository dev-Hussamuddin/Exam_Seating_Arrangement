"""Read the finalized static students/rooms and dynamic daily timetable files."""

from datetime import datetime
from pathlib import Path
import re

from openpyxl import load_workbook

from excel_importer import exam_date, integer, normalized_label, parse_roll_list, parse_time_value, text


FILENAMES = ("Students & Subject Data.xlsx", "Classroom Data.xlsx", "Timetable.xlsx")
STUDENT_HEADERS = ("Department", "Class", "Semester", "Subject", "Roll Nos. (Range / List)", "Excluded Roll Nos.")
ROOM_HEADERS = ("Room No.", "Capacity (No. of Benches/Seats)")
TIMETABLE_HEADERS = ("Class", "Semester", "Subject")
EXAM_LABELS = ("Exam Date", "Exam Name", "Month / Year", "Timing")


def three_file_input_available(directory):
    """Partial finalized inputs are an error, never a silent fallback to legacy."""
    paths = [Path(directory) / name for name in FILENAMES]
    if not any(p.exists() for p in paths):
        return False
    missing = [p.name for p in paths if not p.is_file()]
    if missing:
        raise ValueError("Missing finalized input workbook(s): " + ", ".join(missing))
    return True


def _read_table(path, headers):
    book = load_workbook(path, data_only=False)
    try:
        candidates = []
        expected = [normalized_label(h) for h in headers]
        for sheet in book:
            rows = list(sheet.iter_rows(values_only=True))
            for index, row in enumerate(rows):
                labels = [normalized_label(v) for v in row]
                if all(h in labels for h in expected):
                    if any(labels.count(h) != 1 for h in expected):
                        raise ValueError(f"{path.name}: duplicate table header.")
                    candidates.append((rows, index, [labels.index(h) for h in expected]))
        if len(candidates) != 1:
            raise ValueError(f"{path.name}: exactly one table with headers {', '.join(headers)} is required.")
        rows, header, columns = candidates[0]
        for number, row in enumerate(rows, 1):
            if any(isinstance(v, str) and v.startswith("=") for v in row):
                raise ValueError(f"{path.name}, row {number}: use values, not formulas.")
        data = []
        for number, row in enumerate(rows[header + 1:], header + 2):
            values = [row[i] if i < len(row) else None for i in columns]
            if any(v is not None and str(v).strip() for v in values):
                data.append((number, values))
        if not data:
            raise ValueError(f"{path.name}: the table has no data rows.")
        return rows[:header], data
    finally:
        book.close()


def _key(cls, semester, subject):
    return tuple(normalized_label(value) for value in (cls, semester, subject))


def read_three_file_input(directory):
    directory = Path(directory)
    if not three_file_input_available(directory):
        raise ValueError("Place the three finalized workbooks in " + str(directory))
    records = {"Class Data": [], "Classroom Data": [], "Timetable": []}
    _, students = _read_table(directory / FILENAMES[0], STUDENT_HEADERS)
    groups, departments = {}, {}
    for number, row in students:
        try:
            dept, cls, semester, subject = (text(row[0], "Department", 100), text(row[1], "Class", 50),
                                           text(row[2], "Semester", 50), text(row[3], "Subject", 150))
            class_key = normalized_label(cls)
            if departments.setdefault(class_key, normalized_label(dept)) != normalized_label(dept):
                raise ValueError("Use distinct class names for classes in different departments.")
            rolls = set(parse_roll_list(row[4]))
            excluded_value = None if str(row[5]).strip() in ("-", "–", "—") else row[5]
            excluded = set(parse_roll_list(excluded_value, "Excluded Roll Nos.", allow_empty=True))
            if not excluded <= rolls:
                raise ValueError("Excluded rolls must belong to the supplied roll range/list.")
            widths = [len(token) for token in re.findall(r"\d+", str(row[4])) if len(token) > 1 and token.startswith("0")]
            width = max(widths, default=0)
            key = _key(cls, semester, subject)
            if key in groups:
                group = groups[key]
                if rolls & group["included"]:
                    raise ValueError("Duplicate/overlapping ranges for this class, semester and subject.")
                group["included"].update(rolls)
                group["excluded"].update(excluded)
                group["roll_width"] = max(group["roll_width"], width)
            else:
                groups[key] = dict(row=number, class_name=cls, department=dept, semester=semester,
                                   subject=subject, included=rolls, excluded=excluded, roll_width=width)
        except ValueError as error:
            raise ValueError(f"{FILENAMES[0]}, row {number}: {error}") from None
    for group in groups.values():
        rolls, excluded = sorted(group["included"]), sorted(group["excluded"])
        group["rolls"] = sorted(group["included"] - group["excluded"])
        omitted = text(",".join(map(str, excluded)), "Excluded Roll Nos.", 255, required=False)
        records["Class Data"].append((group["row"], (group["class_name"], group["department"],
            rolls[0], rolls[-1], len(group["rolls"]), omitted, group["subject"],
            ",".join(map(str, rolls)), group["semester"])))

    _, room_rows = _read_table(directory / FILENAMES[1], ROOM_HEADERS)
    seen_rooms = set()
    for number, row in room_rows:
        room = text(row[0], "Room No.", 20)
        if normalized_label(room) in seen_rooms:
            raise ValueError(f"{FILENAMES[1]}, row {number}: duplicate room.")
        seen_rooms.add(normalized_label(room))
        records["Classroom Data"].append((number, (room, integer(row[1], "Capacity", minimum=1))))

    top, timetable = _read_table(directory / FILENAMES[2], TIMETABLE_HEADERS)
    settings = {}
    labels = {normalized_label(label) for label in EXAM_LABELS}
    for row in top:
        if not row or normalized_label(row[0]) not in labels:
            continue
        label = normalized_label(row[0])
        if label in settings:
            raise ValueError(f"Timetable.xlsx: duplicate {row[0]}.")
        settings[label] = row[1] if len(row) > 1 else None
    if set(settings) != labels:
        raise ValueError("Timetable.xlsx: top configuration requires " + ", ".join(EXAM_LABELS))
    raw_date = settings["exam date"]
    if isinstance(raw_date, str) and re.fullmatch(r"\d{2}/\d{2}/\d{4}", raw_date.strip()):
        raw_date = datetime.strptime(raw_date.strip(), "%d/%m/%Y").date()
    day = exam_date(raw_date)
    timing = text(settings["timing"], "Timing", 255, required=False)
    if timing:
        parsed = re.sub(r"\b([ap])\.?\s*m\.?", lambda m: m[1].upper() + "M", timing, flags=re.I)
        _, start, end = parse_time_value(parsed)
        if end is None:
            raise ValueError("Timing must be a complete start-end range or blank.")
    else:
        start = end = None
    info = dict(exam_date=day, exam_name=text(settings["exam name"], "Exam Name", 150),
                month_year=text(settings["month / year"], "Month / Year", 100),
                timing=timing, start_time=start, end_time=end)
    selected, seen, identities = [], set(), set()
    for number, row in timetable:
        cls, semester, subject = (text(row[0], "Class", 50), text(row[1], "Semester", 50), text(row[2], "Subject", 150))
        key = _key(cls, semester, subject)
        if key in seen:
            raise ValueError(f"Timetable.xlsx, row {number}: duplicate exam.")
        seen.add(key)
        if key not in groups:
            raise ValueError(f"Timetable.xlsx, row {number}: no matching student data for {cls} / {semester} / {subject}.")
        group = groups[key]
        students_today = {(normalized_label(group["class_name"]), roll) for roll in group["rolls"]}
        if identities & students_today:
            raise ValueError(f"Timetable.xlsx, row {number}: overlapping student groups in the same slot.")
        identities.update(students_today)
        selected.append(group)
        records["Timetable"].append((number, (group["class_name"], group["subject"], group["semester"], day, start, end, timing)))
    # Validate before connecting to the database; static unscheduled groups do not count.
    eligible = sum(len(group["rolls"]) for group in selected)
    capacity = sum(row[1] for _, row in records["Classroom Data"])
    if eligible > capacity:
        raise ValueError(f"Insufficient capacity: {eligible} students require seats but only {capacity} seats are available.")
    return dict(records=records, exam_info=info, exams=selected)
