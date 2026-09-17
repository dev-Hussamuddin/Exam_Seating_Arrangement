"""Offline Excel output and post-commit integration tests."""

from copy import deepcopy
from datetime import date, datetime, time, timedelta
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import Mock, patch

from openpyxl import load_workbook

import seating_allocator
from seating_report import write_seating_report


def sample():
    exams = [dict(timetable_id=1, class_id=1, class_name="FYCS", subject="Maths",
                  department="CS", rolls=[1, 2, 4]),
             dict(timetable_id=2, class_id=1, class_name="FYCS", subject="Science",
                  department="CS", rolls=[8, 10])]
    rooms = [dict(classroom_id=1, classroom_no="001", capacity=6),
             dict(classroom_id=2, classroom_no="002", capacity=4)]
    rows = seating_allocator.allocate_slot(exams, rooms)
    return [((date(2026, 10, 9), timedelta(hours=9), timedelta(hours=11)), exams, rooms, rows)]


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "new" / "output" / "Seating_Arrangement.xlsx"

    def read_report(self, summaries):
        self.assertEqual(write_seating_report(summaries, self.path), self.path)
        book = load_workbook(self.path)
        self.addCleanup(book.close)
        return book.active

    def test_groups_exclusions_unused_rooms_and_summary(self):
        summaries = sample()
        before = deepcopy(summaries)
        sheet = self.read_report(summaries)
        self.assertEqual(summaries, before)
        self.assertEqual([sheet.cell(r, 2).value for r in range(2, 6)], [5, 5, 5, 5])
        self.assertEqual(list(sheet.values)[8:], [
            (datetime(2026, 10, 9), time(9), time(11), "001", 6, "FYCS", "Maths", "1-2,4", 3, 1),
            (datetime(2026, 10, 9), time(9), time(11), "001", None, "FYCS", "Science", "8,10", 2, None),
            (datetime(2026, 10, 9), time(9), time(11), "002", 4, None, None, None, 0, 4)])

    def test_multiple_slots_count_capacity_once_per_room_per_slot(self):
        summaries = sample()
        second = deepcopy(summaries[0])
        summaries.append((("2026-10-10", "13:30:00", "15:00:00"), *second[1:]))
        sheet = self.read_report(summaries)
        self.assertEqual([sheet.cell(r, 2).value for r in range(2, 6)], [10, 10, 10, 10])
        self.assertEqual(sum(sheet.cell(r, 5).value or 0 for r in range(9, 15)), 20)
        self.assertEqual(sheet["B12"].value, time(13, 30))

    def test_empty_and_zero_eligible(self):
        sheet = self.read_report([])
        self.assertEqual(sheet.max_row, 8)
        self.assertEqual([sheet.cell(r, 2).value for r in range(2, 6)], [0, 0, 0, 0])
        slot, _, rooms, _ = sample()[0]
        sheet = self.read_report([(slot, [], rooms, [])])
        self.assertEqual(sheet["B5"].value, 10)
        self.assertEqual(sheet["I9"].value, 0)

    def test_formatting_and_literal_subject(self):
        summaries = sample()
        summaries[0][1][0]["subject"] = "=1+1"
        sheet = self.read_report(summaries)
        self.assertEqual(sheet["G9"].data_type, "s")
        self.assertEqual(sheet["G9"].value, "=1+1")
        self.assertTrue(sheet["A8"].font.bold)
        self.assertEqual(sheet["H9"].border.bottom.style, "thin")
        self.assertEqual(sheet["A9"].number_format, "dd-mmm-yyyy")
        self.assertEqual(sheet["B9"].number_format, "hh:mm")
        self.assertEqual(sheet.freeze_panes, "F9")
        self.assertGreaterEqual(sheet.column_dimensions["H"].width, 50)

    def test_repeat_export_replaces_previous_report(self):
        write_seating_report(sample(), self.path)
        sheet = self.read_report([])
        self.assertEqual(sheet.max_row, 8)
        self.assertEqual(list(self.path.parent.iterdir()), [self.path])

    def test_failed_save_preserves_previous_file_and_cleans_temporary(self):
        write_seating_report(sample(), self.path)
        before = self.path.read_bytes()
        with patch("seating_report.Workbook.save", side_effect=OSError("disk full")):
            with self.assertRaises(OSError):
                write_seating_report([], self.path)
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual(list(self.path.parent.iterdir()), [self.path])

    def test_default_output_location(self):
        with patch("seating_report.OUTPUT_DIR", Path(self.temp.name)):
            path = write_seating_report([])
        self.assertEqual(path, Path(self.temp.name) / "Seating_Arrangement.xlsx")
        self.assertTrue(path.is_file())


class ReportIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.connection = Mock()
        for name in ("migrate_input_schema", "migrate_seating_schema", "print_summary"):
            patcher = patch.object(seating_allocator, name)
            patcher.start()
            self.addCleanup(patcher.stop)
        patcher = patch.object(seating_allocator, "get_connection", return_value=self.connection)
        patcher.start()
        self.addCleanup(patcher.stop)
        patcher = patch("builtins.print")
        self.print = patcher.start()
        self.addCleanup(patcher.stop)

    def test_report_runs_after_successful_allocation(self):
        summaries = sample()
        events = []
        def allocate(connection):
            events.append("allocation committed")
            return summaries
        def report(data):
            self.assertIs(data, summaries)
            events.append("report")
            return Path("output/Seating_Arrangement.xlsx")
        with patch.object(seating_allocator, "generate_arrangements", side_effect=allocate), \
             patch.object(seating_allocator, "write_seating_report", side_effect=report):
            self.assertEqual(seating_allocator.main(), 0)
        self.assertEqual(events, ["allocation committed", "report"])
        self.connection.close.assert_called_once()

    def test_no_report_when_allocation_fails_or_no_slots(self):
        for outcome in (ValueError("Insufficient capacity"), []):
            with self.subTest(outcome=outcome), \
                 patch.object(seating_allocator, "generate_arrangements") as allocate, \
                 patch.object(seating_allocator, "write_seating_report") as report:
                if isinstance(outcome, Exception):
                    allocate.side_effect = outcome
                else:
                    allocate.return_value = outcome
                self.assertEqual(seating_allocator.main(), 1 if isinstance(outcome, Exception) else 0)
                report.assert_not_called()

    def test_export_failure_reports_committed_allocation_without_rollback(self):
        with patch.object(seating_allocator, "generate_arrangements", return_value=sample()), \
             patch.object(seating_allocator, "write_seating_report", side_effect=PermissionError("file is open")):
            self.assertEqual(seating_allocator.main(), 1)
        self.connection.rollback.assert_not_called()
        self.connection.close.assert_called_once()
        self.assertTrue(any("Allocation succeeded and was committed" in str(c) for c in self.print.call_args_list))


if __name__ == "__main__":
    unittest.main()
