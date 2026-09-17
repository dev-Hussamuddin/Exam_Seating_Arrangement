"""Exact college layout, merged classes, subject roll lists and migration checks."""

import unittest
import io
from contextlib import redirect_stdout
from datetime import time, date, datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock, patch

from openpyxl import Workbook

import excel_importer as importer
from database import migrate_input_schema
from seating_allocator import generate_arrangements, load_slot
from test_excel_importer import TestConnection


class CollegeInputTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "Tables.xlsx"
        self.connection = TestConnection()
        self.addCleanup(self.connection.db.close)
        self.workbook = Workbook()
        self.addCleanup(self.workbook.close)
        self.sheet = self.workbook.active
        self.sheet.title = "Original college sheet"
        for name, column in (("Class Data", 2), ("Classroom Data", 9), ("Timetable", 13)):
            self.sheet.cell(3, column, name)
            for c, value in enumerate(importer.COLLEGE_HEADERS[name], column):
                self.sheet.cell(5, c, value)
        for r, row in enumerate([
            ["SYCS", "Maths", "1-4,8,10-12", "2,10", 6],
            [None, "Science", "21,23-25", "24", 3],
            ["SYBSc", "English", "1-2", None, 2],
        ], 6):
            for c, value in enumerate(row, 2):
                self.sheet.cell(r, c, value)
        self.sheet.merge_cells("B6:B7")
        self.sheet.cell(6, 9, "308")
        self.sheet.cell(6, 10, 26)
        for r, row in enumerate([
            ["SYCS", "Maths", "09:00-10:00", date(2026, 10, 9)],
            [None, "Science", "11:00 AM - 12:00 PM", "2026-10-10"],
            ["SYBSc", "English", "09:00", datetime(2026, 10, 11)],
        ], 6):
            for c, value in enumerate(row, 13):
                self.sheet.cell(r, c, value)
        self.sheet.merge_cells("M6:M7")

    def run_import(self):
        self.workbook.save(self.path)
        with patch.object(importer, "get_connection", return_value=self.connection), patch.object(importer, "migrate_input_schema"):
            return importer.import_workbook(self.path)

    def test_exact_layout_and_subject_lists(self):
        counts = self.run_import()
        self.assertEqual(counts, dict(classes=2, batches=3, classrooms=1, timetable=3))
        db = self.connection.db
        self.assertEqual(db.execute("SELECT subject, roll_numbers, non_included_rolls FROM student_batches ORDER BY batch_id").fetchall(),
                         [("Maths", "1,2,3,4,8,10,11,12", "2,10"), ("Science", "21,23,24,25", "24"), ("English", "1,2", None)])
        self.assertEqual(db.execute("SELECT COUNT(*) FROM timetable WHERE exam_date IS NULL").fetchone()[0], 0)
        self.assertEqual(db.execute("SELECT time_text FROM timetable WHERE subject='English'").fetchone()[0], "09:00")

    def test_repeat_import(self):
        self.run_import()
        self.sheet.cell(7, 4, "21,23-26")
        counts = self.run_import()
        self.assertEqual(counts["classes"], 0)
        self.assertEqual(counts["timetable"], 0)
        self.assertEqual(self.connection.db.execute("SELECT COUNT(*) FROM student_batches").fetchone()[0], 3)

    def test_date_column_detected_after_move(self):
        for row in range(5, 9):
            self.sheet.cell(row, 19, self.sheet.cell(row, 16).value)
            self.sheet.cell(row, 16).value = None
        self.run_import()
        self.assertEqual(self.connection.db.execute("SELECT exam_date FROM timetable ORDER BY timetable_id").fetchall(),
                         [("2026-10-09",), ("2026-10-10",), ("2026-10-11",)])

    def test_real_section_heading_aliases(self):
        self.sheet["B3"] = "Class Table"
        self.sheet["I3"] = "Classroom Table"
        self.sheet["M3"] = "Time table"
        self.assertEqual(self.run_import()["timetable"], 3)
        self.assertEqual(self.connection.db.execute("SELECT exam_date FROM timetable ORDER BY timetable_id").fetchall(),
                         [("2026-10-09",), ("2026-10-10",), ("2026-10-11",)])

    def test_heading_alias_case_and_spacing(self):
        self.sheet["B3"] = " CLASS  TABLE "
        self.sheet["I3"] = "classroom\nTable"
        self.sheet["M3"] = "Time   TABLE"
        self.assertEqual(self.run_import()["timetable"], 3)

    def test_missing_or_invalid_date_rejected_before_database(self):
        for value in (None, "", "2026-02-30", "09/10/2026", "2026-1-9", 46000):
            with self.subTest(value=value):
                self.sheet.cell(6, 16).value = value
                self.sheet.cell(6, 16).number_format = "General"
                with self.assertRaisesRegex(ValueError, "Exam Date"):
                    self.run_import()
        self.assertEqual(self.connection.db.execute("SELECT COUNT(*) FROM timetable").fetchone()[0], 0)

    def test_missing_date_header_rejected(self):
        self.sheet.cell(5, 16).value = None
        with self.assertRaisesRegex(ValueError, "Exam Date"):
            self.run_import()

    def test_null_date_updated_in_place_and_rerun_is_idempotent(self):
        self.run_import()
        db = self.connection.db
        ids = db.execute("SELECT timetable_id FROM timetable ORDER BY timetable_id").fetchall()
        db.execute("UPDATE timetable SET exam_date=NULL, time_text='old formatting'")
        db.commit()
        self.assertEqual(self.run_import()["timetable"], 3)
        self.assertEqual(self.run_import()["timetable"], 0)
        self.assertEqual(db.execute("SELECT timetable_id FROM timetable ORDER BY timetable_id").fetchall(), ids)
        self.assertEqual(db.execute("SELECT COUNT(*) FROM timetable WHERE exam_date IS NULL").fetchone()[0], 0)

    def test_date_repair_rolls_back_if_later_match_is_ambiguous(self):
        self.run_import()
        db = self.connection.db
        db.execute("UPDATE timetable SET exam_date=NULL")
        db.execute("""INSERT INTO timetable (class_id,subject,semester,exam_date,start_time,end_time,time_text)
            SELECT class_id,subject,semester,exam_date,start_time,end_time,time_text FROM timetable WHERE subject='Science'""")
        db.commit()
        before = list(db.iterdump())
        with self.assertRaisesRegex(ValueError, "multiple matching undated"):
            self.run_import()
        self.assertEqual(list(db.iterdump()), before)

    def test_other_time_slot_not_repaired(self):
        self.run_import()
        db = self.connection.db
        db.execute("UPDATE timetable SET exam_date=NULL, start_time='08:00:00' WHERE subject='Maths'")
        db.commit()
        self.assertEqual(self.run_import()["timetable"], 1)
        self.assertEqual(db.execute("SELECT COUNT(*) FROM timetable WHERE subject='Maths'").fetchone()[0], 2)

    def test_human_readable_time_ranges(self):
        for raw, start, end in (
            ("12:00-01:00", time(12), time(13)),
            ("11:30-01:00", time(11, 30), time(13)),
            ("13:00-02:00", time(13), time(14)),
            ("12:00\u201301:00", time(12), time(13)),
            ("12:00 PM to 1:00 PM", time(12), time(13)),
            ("12:00 AM-01:00 AM", time(0), time(1)),
            ("09:00-10:00", time(9), time(10)),
            ("12:00:00-01:30:00", time(12), time(13, 30)),
        ):
            with self.subTest(raw=raw):
                self.assertEqual(importer.parse_time_value(raw), (raw, start, end))

    def test_excel_time_cell_and_single_time(self):
        self.assertEqual(importer.parse_time_value(time(12, 30)), ("12:30:00", time(12, 30), None))
        self.assertEqual(importer.parse_time_value("09:00"), ("09:00", time(9), None))
        self.sheet.cell(6, 15, time(12, 30))
        self.run_import()
        self.assertEqual(self.connection.db.execute("SELECT start_time, end_time FROM timetable WHERE subject='Maths'").fetchone(),
                         ("12:30:00", None))

    def test_invalid_or_explicitly_reversed_times(self):
        for raw in ("12:00-12:00", "12:00 PM-01:00 AM", "23:00-01:00",
                    "25:00-26:00", "12:99-01:00", "bad", "12:00-", "09:00-10:00-11:00"):
            with self.subTest(raw=raw), self.assertRaises(ValueError):
                importer.parse_time_value(raw)

    def test_noon_range_repeat_import_preserves_source(self):
        self.sheet.cell(6, 15, "12:00-01:00")
        self.run_import()
        self.assertEqual(self.run_import()["timetable"], 0)
        self.assertEqual(self.connection.db.execute("SELECT start_time, end_time, time_text FROM timetable WHERE subject='Maths'").fetchone(),
                         ("12:00:00", "13:00:00", "12:00-01:00"))

    def test_actual_trailing_separators_and_unicode_ranges(self):
        self.assertEqual(importer.parse_roll_list("1 - 63,"), list(range(1, 64)))
        self.assertEqual(importer.parse_roll_list("22, 46-49, "), [22, 46, 47, 48, 49])
        self.assertEqual(importer.parse_roll_list("1\u20132, 15\u201360, 62\u201363"),
                         [1, 2] + list(range(15, 61)) + [62, 63])
        for invalid in (None, "", ",", "1,,3", "4-1", "1-x"):
            with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                importer.parse_roll_list(invalid)

    def test_actual_merged_subject_group_and_blank_rolls(self):
        self.sheet.cell(6, 4, "1 - 63,")
        self.sheet.cell(6, 5, 58)
        self.sheet.cell(6, 6, 62)
        for col in (4, 5, 6):
            self.sheet.merge_cells(start_row=6, end_row=7, start_column=col, end_column=col)
        self.run_import()
        batches = self.connection.db.execute("SELECT roll_numbers, non_included_rolls, number_of_students FROM student_batches WHERE class_id=1").fetchall()
        self.assertEqual(len(batches), 2)
        self.assertEqual(batches[0], batches[1])
        self.assertEqual(batches[0], (",".join(map(str, range(1, 64))), "58", 62))

    def test_merged_subject_continuation_is_not_another_record(self):
        # Reproduce Sports-style merged Subject, Rolls and Count across blank rows.
        for col in range(2, 7):
            self.sheet.merge_cells(start_row=8, end_row=10, start_column=col, end_column=col)
        self.assertEqual(self.run_import()["batches"], 3)

    def test_unrelated_blank_rolls_do_not_inherit_previous_subject(self):
        self.sheet.cell(8, 4).value = None
        with self.assertRaisesRegex(ValueError, "Roll nos. is required"):
            self.run_import()

    def test_centered_title_and_staggered_merged_headers(self):
        self.sheet["B3"] = None
        self.sheet["C3"] = "Class Data"
        self.sheet.merge_cells("C3:D3")
        for col in (2, 4, 6):
            self.sheet.cell(4, col, self.sheet.cell(5, col).value)
            self.sheet.cell(5, col).value = None
            self.sheet.merge_cells(start_row=4, end_row=5, start_column=col, end_column=col)
        self.assertEqual(self.run_import()["batches"], 3)

    def test_merged_rolls_and_counts_shared_by_subjects(self):
        for col in (4, 5, 6):
            self.sheet.merge_cells(start_row=6, end_row=7, start_column=col, end_column=col)
        self.run_import()
        records = self.connection.db.execute("SELECT roll_numbers, non_included_rolls, number_of_students FROM student_batches WHERE class_id=1").fetchall()
        self.assertEqual(records[0], records[1])

    def test_subject_specific_allocator_loading(self):
        self.run_import()
        db = self.connection.db
        db.execute("UPDATE timetable SET exam_date='2026-10-09' WHERE subject='Maths'")
        db.commit()
        cursor = self.connection.cursor()
        records = load_slot(cursor, ("2026-10-09", "09:00:00", "10:00:00"))
        cursor.close()
        self.assertEqual(records[0]["rolls"], [1, 3, 4, 8, 11, 12])

    def test_undated_allocation_stops(self):
        self.run_import()
        with self.assertRaisesRegex(ValueError, "undated"):
            generate_arrangements(self.connection)

    def test_unknown_subject_is_preserved(self):
        self.sheet.cell(7, 14, "Missing subject")
        output = io.StringIO()
        with redirect_stdout(output):
            self.run_import()
        self.assertIn("skipped for seating", output.getvalue())
        self.assertEqual(self.connection.db.execute("SELECT COUNT(*) FROM timetable").fetchone()[0], 3)

    def test_timetable_only_class_preserved_without_students(self):
        self.sheet.unmerge_cells("M6:M7")
        self.sheet.cell(6, 13, "FYCS")
        self.sheet.cell(7, 13, "SYCS")
        output = io.StringIO()
        with redirect_stdout(output):
            self.run_import()
            repeated = self.run_import()
        db = self.connection.db
        self.assertIn("FYCS / Maths: no student data", output.getvalue())
        self.assertEqual(repeated["timetable"], 0)
        self.assertEqual(db.execute("SELECT COUNT(*) FROM student_batches b JOIN classes c ON c.class_id=b.class_id WHERE c.class_name='FYCS'").fetchone()[0], 0)
        self.assertEqual(db.execute("SELECT COUNT(*) FROM timetable t JOIN classes c ON c.class_id=t.class_id WHERE c.class_name='FYCS'").fetchone()[0], 1)
        self.assertEqual(db.execute("SELECT COUNT(*) FROM student_batches").fetchone()[0], 3)

    def test_exclusion_outside_subject_list(self):
        self.sheet.cell(6, 5, "7")
        with self.assertRaisesRegex(ValueError, "subject's roll list"):
            self.run_import()

    def test_existing_department_preserved(self):
        db = self.connection.db
        db.execute("INSERT INTO classes VALUES (1, 'SYCS', 'CS')")
        db.commit()
        self.run_import()
        self.assertEqual(db.execute("SELECT department FROM classes WHERE class_name='SYCS'").fetchone()[0], "CS")

    def test_unmerged_continuation(self):
        self.sheet.unmerge_cells("B6:B7")
        self.sheet.unmerge_cells("M6:M7")
        self.assertEqual(self.run_import()["batches"], 3)

    def test_blank_separator_does_not_carry_class(self):
        self.sheet.unmerge_cells("B6:B7")
        for col in range(2, 7):
            self.sheet.cell(7, col).value = None
        self.sheet.cell(8, 2).value = None
        with self.assertRaisesRegex(ValueError, "Class is required"):
            self.run_import()

    def test_migration_is_additive_and_repeatable(self):
        connection = Mock()
        cursor = connection.cursor.return_value
        cursor.fetchone.side_effect = [None, None, None, ("exam_date", "date", "NO"),
                                       ("start_time", "time", "NO"), ("end_time", "time", "NO")]
        migrate_input_schema(connection)
        statements = [call.args[0] for call in cursor.execute.call_args_list]
        self.assertEqual(sum(s.startswith("ALTER") for s in statements), 6)
        self.assertFalse(any("DROP" in s or "DELETE" in s for s in statements))
        cursor.reset_mock()
        cursor.fetchone.side_effect = [("exists", "text", "YES")] * 6
        migrate_input_schema(connection)
        self.assertTrue(all(call.args[0].startswith("SHOW") for call in cursor.execute.call_args_list))


if __name__ == "__main__":
    unittest.main()
